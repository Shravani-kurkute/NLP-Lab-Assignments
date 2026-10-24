#!/usr/bin/env python3
from pathlib import Path
import json

import joblib
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from sklearn.metrics import (
    accuracy_score,
    classification_report,
    confusion_matrix,
    precision_recall_fscore_support,
)


# PATH CONFIGURATION
BASE_DIR = Path(__file__).resolve().parent

ARTIFACT_DIR = BASE_DIR / "artifacts"
VISUALIZATION_DIR = ARTIFACT_DIR / "visualization"

MODEL_PATH = ARTIFACT_DIR / "news_classifier.joblib"
TEST_DATA_PATH = ARTIFACT_DIR / "test_data.csv"
METADATA_PATH = ARTIFACT_DIR / "training_metadata.json"

RESULTS_DIR = BASE_DIR / "results"
EXCEL_PATH = RESULTS_DIR / "News_Article_Classification_Test_Results.xlsx"
CONFUSION_MATRIX_PATH = RESULTS_DIR / "confusion_matrix.png"
CLASSIFICATION_REPORT_PATH = RESULTS_DIR / "classification_report_visualization.png"
METRICS_CHART_PATH = RESULTS_DIR / "per_category_metrics.png"


def plot_confusion_matrix(y_true, y_pred, classes, output_path: Path) -> None:
    cm = confusion_matrix(y_true, y_pred, labels=classes)
    
    fig, ax = plt.subplots(figsize=(10, 8))
    sns.heatmap(
        cm,
        annot=True,
        fmt="d",
        cmap="Blues",
        xticklabels=classes,
        yticklabels=classes,
        cbar=True,
        ax=ax,
    )
    
    ax.set_title("Confusion Matrix - News Article Classification", fontsize=16, fontweight="bold")
    ax.set_ylabel("Actual Category", fontsize=12)
    ax.set_xlabel("Predicted Category", fontsize=12)
    plt.xticks(rotation=45, ha="right")
    plt.yticks(rotation=0)
    
    plt.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    print(f"Saved confusion matrix: {output_path}")


def plot_classification_report_visualization(report_dict, classes, output_path: Path) -> None:
    metrics = []
    for cls in classes:
        metrics.append({
            "Category": cls,
            "Precision": report_dict[cls]["precision"],
            "Recall": report_dict[cls]["recall"],
            "F1-Score": report_dict[cls]["f1-score"],
            "Support": report_dict[cls]["support"],
        })
    
    df_metrics = pd.DataFrame(metrics)
    
    fig, axes = plt.subplots(2, 2, figsize=(16, 12))
    
    # Precision
    bars1 = axes[0, 0].bar(df_metrics["Category"], df_metrics["Precision"], color="skyblue")
    axes[0, 0].set_title("Precision by Category", fontsize=14, fontweight="bold")
    axes[0, 0].set_ylabel("Precision")
    axes[0, 0].set_ylim(0, 1.1)
    axes[0, 0].bar_label(bars1, fmt="%.3f", padding=3)
    axes[0, 0].tick_params(axis="x", rotation=45)
    axes[0, 0].grid(axis="y", linestyle="--", alpha=0.4)
    
    # Recall
    bars2 = axes[0, 1].bar(df_metrics["Category"], df_metrics["Recall"], color="lightgreen")
    axes[0, 1].set_title("Recall by Category", fontsize=14, fontweight="bold")
    axes[0, 1].set_ylabel("Recall")
    axes[0, 1].set_ylim(0, 1.1)
    axes[0, 1].bar_label(bars2, fmt="%.3f", padding=3)
    axes[0, 1].tick_params(axis="x", rotation=45)
    axes[0, 1].grid(axis="y", linestyle="--", alpha=0.4)
    
    # F1-Score
    bars3 = axes[1, 0].bar(df_metrics["Category"], df_metrics["F1-Score"], color="coral")
    axes[1, 0].set_title("F1-Score by Category", fontsize=14, fontweight="bold")
    axes[1, 0].set_ylabel("F1-Score")
    axes[1, 0].set_ylim(0, 1.1)
    axes[1, 0].bar_label(bars3, fmt="%.3f", padding=3)
    axes[1, 0].tick_params(axis="x", rotation=45)
    axes[1, 0].grid(axis="y", linestyle="--", alpha=0.4)
    
    # Support
    bars4 = axes[1, 1].bar(df_metrics["Category"], df_metrics["Support"], color="gold")
    axes[1, 1].set_title("Support (Number of Test Samples) by Category", fontsize=14, fontweight="bold")
    axes[1, 1].set_ylabel("Count")
    axes[1, 1].bar_label(bars4, fmt="%d", padding=3)
    axes[1, 1].tick_params(axis="x", rotation=45)
    axes[1, 1].grid(axis="y", linestyle="--", alpha=0.4)
    
    plt.suptitle("Classification Report Visualization", fontsize=18, fontweight="bold", y=0.995)
    plt.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    print(f"Saved classification report visualization: {output_path}")


def plot_per_category_metrics(report_dict, classes, output_path: Path) -> None:
    categories = classes
    precision = [report_dict[c]["precision"] for c in categories]
    recall = [report_dict[c]["recall"] for c in categories]
    f1 = [report_dict[c]["f1-score"] for c in categories]
    
    x = np.arange(len(categories))
    width = 0.25
    
    fig, ax = plt.subplots(figsize=(12, 7))
    bars1 = ax.bar(x - width, precision, width, label="Precision", color="skyblue")
    bars2 = ax.bar(x, recall, width, label="Recall", color="lightgreen")
    bars3 = ax.bar(x + width, f1, width, label="F1-Score", color="coral")
    
    ax.set_title("Per-Category Classification Metrics", fontsize=16, fontweight="bold")
    ax.set_ylabel("Score")
    ax.set_xlabel("News Category")
    ax.set_xticks(x)
    ax.set_xticklabels(categories, rotation=45, ha="right")
    ax.legend()
    ax.set_ylim(0, 1.15)
    ax.grid(axis="y", linestyle="--", alpha=0.4)
    
    ax.bar_label(bars1, fmt="%.3f", padding=2, fontsize=8)
    ax.bar_label(bars2, fmt="%.3f", padding=2, fontsize=8)
    ax.bar_label(bars3, fmt="%.3f", padding=2, fontsize=8)
    
    plt.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    print(f"Saved per-category metrics chart: {output_path}")


def create_excel_report(test_results_df, report_dict, overall_metrics, output_path: Path) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        # Sheet 1: Test Results
        test_results_df.to_excel(writer, sheet_name="Test Results", index=False)
        
        # Sheet 2: Metrics
        metrics_data = []
        metrics_data.append({"Metric": "Accuracy", "Value": overall_metrics["accuracy"]})
        metrics_data.append({"Metric": "Macro Precision", "Value": overall_metrics["macro_precision"]})
        metrics_data.append({"Metric": "Macro Recall", "Value": overall_metrics["macro_recall"]})
        metrics_data.append({"Metric": "Macro F1-Score", "Value": overall_metrics["macro_f1"]})
        metrics_data.append({"Metric": "Weighted Precision", "Value": overall_metrics["weighted_precision"]})
        metrics_data.append({"Metric": "Weighted Recall", "Value": overall_metrics["weighted_recall"]})
        metrics_data.append({"Metric": "Weighted F1-Score", "Value": overall_metrics["weighted_f1"]})
        metrics_data.append({"Metric": "Correct Predictions", "Value": overall_metrics["correct_count"]})
        metrics_data.append({"Metric": "Incorrect Predictions", "Value": overall_metrics["incorrect_count"]})
        metrics_data.append({"Metric": "Total Test Samples", "Value": overall_metrics["total_samples"]})
        
        metrics_df = pd.DataFrame(metrics_data)
        metrics_df.to_excel(writer, sheet_name="Metrics", index=False)
        
        # Sheet 3: Classification Report
        report_data = []
        for key, value in report_dict.items():
            if key in ["accuracy"]:
                continue
            if isinstance(value, dict):
                report_data.append({
                    "Category": key,
                    "Precision": value["precision"],
                    "Recall": value["recall"],
                    "F1-Score": value["f1-score"],
                    "Support": value["support"],
                })
            elif key in ["macro avg", "weighted avg"]:
                report_data.append({
                    "Category": key,
                    "Precision": value["precision"],
                    "Recall": value["recall"],
                    "F1-Score": value["f1-score"],
                    "Support": value["support"],
                })
        
        report_df = pd.DataFrame(report_data)
        report_df.to_excel(writer, sheet_name="Classification Report", index=False)
    
    # Format the workbook
    from openpyxl import load_workbook
    
    wb = load_workbook(output_path)
    
    # Format all sheets
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        # Header formatting
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        
        # Freeze header
        ws.freeze_panes = "A2"
        
        # Add filters
        ws.auto_filter.ref = ws.dimensions
        
        # Adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            if sheet_name == "Test Results" and column_letter == "A":
                adjusted_width = 60  # Text column wider
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Wrap text for Text column
        if sheet_name == "Test Results":
            for row in range(2, ws.max_row + 1):
                ws[f"A{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    
    wb.save(output_path)
    print(f"Saved Excel report: {output_path}")


def main() -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    VISUALIZATION_DIR.mkdir(parents=True, exist_ok=True)
    
    print("Loading saved model and test data...")
    if not MODEL_PATH.exists():
        raise FileNotFoundError(f"Model not found: {MODEL_PATH}. Run train_model.py first.")
    if not TEST_DATA_PATH.exists():
        raise FileNotFoundError(f"Test data not found: {TEST_DATA_PATH}. Run train_model.py first.")
    
    pipeline = joblib.load(MODEL_PATH)
    test_df = pd.read_csv(TEST_DATA_PATH)
    
    print(f"Loaded test data with {len(test_df)} samples")
    
    X_test = test_df["text"]
    y_true = test_df["category"]
    
    print("\nMaking predictions...")
    y_pred = pipeline.predict(X_test)
    
    # Calculate metrics
    accuracy = accuracy_score(y_true, y_pred)
    precision_macro, recall_macro, f1_macro, _ = precision_recall_fscore_support(
        y_true, y_pred, average="macro"
    )
    precision_weighted, recall_weighted, f1_weighted, _ = precision_recall_fscore_support(
        y_true, y_pred, average="weighted"
    )
    
    # Per-category metrics
    report = classification_report(y_true, y_pred, output_dict=True)
    
    # Get classes
    classes = sorted(y_true.unique())
    
    # Count correct/incorrect
    correct_mask = y_true == y_pred
    correct_count = int(correct_mask.sum())
    incorrect_count = int((~correct_mask).sum())
    
    overall_metrics = {
        "accuracy": accuracy,
        "macro_precision": precision_macro,
        "macro_recall": recall_macro,
        "macro_f1": f1_macro,
        "weighted_precision": precision_weighted,
        "weighted_recall": recall_weighted,
        "weighted_f1": f1_weighted,
        "correct_count": correct_count,
        "incorrect_count": incorrect_count,
        "total_samples": len(test_df),
    }
    
    print("\n=== Evaluation Metrics ===")
    print(f"Accuracy: {accuracy:.4f} ({correct_count}/{len(test_df)})")
    print(f"Macro Precision: {precision_macro:.4f}")
    print(f"Macro Recall: {recall_macro:.4f}")
    print(f"Macro F1-Score: {f1_macro:.4f}")
    print(f"Weighted Precision: {precision_weighted:.4f}")
    print(f"Weighted Recall: {recall_weighted:.4f}")
    print(f"Weighted F1-Score: {f1_weighted:.4f}")
    print(f"Correct: {correct_count}, Incorrect: {incorrect_count}")
    
    print("\n=== Classification Report ===")
    print(classification_report(y_true, y_pred))
    
    # Generate visualizations
    print("\nGenerating visualizations...")
    plot_confusion_matrix(y_true, y_pred, classes, CONFUSION_MATRIX_PATH)
    plot_classification_report_visualization(report, classes, CLASSIFICATION_REPORT_PATH)
    plot_per_category_metrics(report, classes, METRICS_CHART_PATH)
    
    # Create test results dataframe
    test_results_df = pd.DataFrame({
        "Text": X_test.values,
        "Actual Category": y_true.values,
        "Predicted Category": y_pred,
        "Result (Pass/Fail)": ["Pass" if actual == pred else "Fail" 
                                for actual, pred in zip(y_true.values, y_pred)],
    })
    
    # Create Excel report
    print("\nCreating Excel report...")
    create_excel_report(test_results_df, report, overall_metrics, EXCEL_PATH)
    
    # Save evaluation summary
    eval_summary = {
        "overall_metrics": overall_metrics,
        "classification_report": report,
        "confusion_matrix": confusion_matrix(y_true, y_pred, labels=classes).tolist(),
        "classes": classes,
    }
    
    eval_summary_path = RESULTS_DIR / "evaluation_summary.json"
    with open(eval_summary_path, "w", encoding="utf-8") as f:
        json.dump(eval_summary, f, indent=4)
    print(f"Saved evaluation summary: {eval_summary_path}")
    
    print("\nTesting completed successfully!")


if __name__ == "__main__":
    main()
