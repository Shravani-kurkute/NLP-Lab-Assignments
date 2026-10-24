#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
NLP Laboratory - Assignment 05
Generate the academic report as DOCX and convert it to PDF with LibreOffice.

Run after PART_A.py and PART_B.py:
    python3 tools/generate_report.py
"""

from __future__ import annotations

import json
import os
import subprocess
import sys

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RESULTS = os.path.join(HERE, "results")
FIGS = os.path.join(HERE, "visualizations")
REPORT_DIR = os.path.join(HERE, "report")

STUDENT = {
    "name": "Shravani Balasaheb Kurkute",
    "roll": "421129",
    "prn": "22311926",
    "cls": "BTechA",
    "year": "2026-2027",
    "faculty": "Dr. Kapil N. Vhatkar",
}
DOCX = os.path.join(REPORT_DIR, f"{STUDENT['roll']}_{STUDENT['name']}_NLP_Assignment05.docx")


def add_table(doc, headers, rows, style="Light Grid Accent 1", font=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = style
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(h)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = "" if val is None else str(val)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font)
    return table


def add_image(doc, filename, width=6.2, caption=None):
    path = os.path.join(FIGS, filename)
    if not os.path.isfile(path):
        return
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        c = doc.add_paragraph(caption)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in c.runs:
            r.italic = True
            r.font.size = Pt(9)


def main() -> None:
    os.makedirs(REPORT_DIR, exist_ok=True)
    a = json.load(open(os.path.join(RESULTS, "part_a_metrics.json")))
    b = json.load(open(os.path.join(RESULTS, "part_b_metrics.json")))
    tm = a["test_metrics"]

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ------------------------------------------------------------- title page
    t = doc.add_heading("NATURAL LANGUAGE PROCESSING", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("LABORATORY ASSIGNMENT – 05")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(16)
    p = doc.add_paragraph("Sentiment Analysis and Mini Project – Research Paper Study & Implementation")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True

    add_table(doc, ["Field", "Details"], [
        ["Student Name", STUDENT["name"]],
        ["Roll No.", STUDENT["roll"]],
        ["PRN / Registration No.", STUDENT["prn"]],
        ["Class / Division", STUDENT["cls"]],
        ["Academic Year", STUDENT["year"]],
        ["Faculty Name", STUDENT["faculty"]],
    ])
    doc.add_paragraph()

    doc.add_heading("Assignment 5: Sentiment Analysis and Mini Project", level=1)
    p = doc.add_paragraph()
    p.add_run("Aim: ").bold = True
    p.add_run(
        "To implement sentiment analysis over a suitable dataset using preprocessing, feature "
        "extraction and a machine-learning algorithm, and to study and implement a research paper "
        "in multidisciplinary NLP using open-source tools and libraries.")

    doc.add_heading("1. Objectives", level=1)
    for o in [
        "Understand the complete workflow of a sentiment-analysis system.",
        "Preprocess textual data and prepare it for machine learning.",
        "Extract useful textual features using Bag-of-Words / TF-IDF.",
        "Train and evaluate suitable machine-learning classifiers.",
        "Interpret evaluation results using standard classification metrics.",
        "Study a research paper in multidisciplinary NLP and implement a reproducible component "
        "of the proposed approach using open-source resources.",
    ]:
        doc.add_paragraph(o, style="List Bullet")

    doc.add_heading("2. Problem Statement", level=1)
    doc.add_paragraph(
        "Select a suitable labelled sentiment dataset (product reviews, movie reviews, social-media "
        "text, etc.). Develop an end-to-end sentiment-classification system involving preprocessing, "
        "feature extraction, model training and evaluation. In addition, select a peer-reviewed "
        "multidisciplinary NLP research paper and implement a meaningful component or a simplified "
        "reproducible version of the proposed approach.")

    # ------------------------------------------------------------------ PART A
    doc.add_heading("3. Part A – Sentiment Analysis", level=1)
    doc.add_heading("3.1 Dataset Details", level=2)
    ds = a["dataset"]
    add_table(doc, ["Item", "Details"], [
        ["Reference dataset", "Amazon Product Reviews Dataset (Kaggle: yasserh/amazon-product-reviews-dataset)"],
        ["File used in this run", ds["file"]],
        ["Source actually used", ds["source"]],
        ["Task / Language", "Sentiment classification / English"],
        ["Raw records", f"{ds['n_raw']}"],
        ["After cleaning", f"{ds['n_clean']}"],
        ["Classes", "2: positive (rating ≥ 4), negative (rating ≤ 2)"],
        ["Training samples", f"{a['split']['n_train']} (80%, stratified)"],
        ["Testing samples", f"{a['split']['n_test']} (20%, stratified)"],
    ])
    doc.add_paragraph(
        "Note on availability: the exact Kaggle file 7817_1.csv was not present locally and Kaggle "
        "downloads require an authenticated account. The pipeline therefore uses a genuine, openly "
        "available Datafiniti Amazon product-review CSV with the identical schema "
        "(name, reviews.title, reviews.text, reviews.rating). To reproduce with the exact reference "
        "file, download 7817_1.csv from the Kaggle link above and place it at data/7817_1.csv; "
        "PART_A.py will then prefer it automatically. No data or results in this report were fabricated.")

    doc.add_heading("3.2 Theory / Background", level=2)
    doc.add_paragraph(
        "Sentiment analysis identifies the opinion or polarity expressed in text. A supervised "
        "pipeline has four stages: data collection, preprocessing, feature extraction, model training "
        "and evaluation. Common representations are Bag-of-Words, TF-IDF, word embeddings and n-grams; "
        "common algorithms include Naive Bayes, Logistic Regression, SVM and tree ensembles. This work "
        "uses TF-IDF (Term Frequency × Inverse Document Frequency) features, which give high weight to "
        "words that are frequent in a review but rare across the corpus.")

    doc.add_heading("3.3 Preprocessing", level=2)
    for step in [
        "Inspect the dataset and identify missing/duplicate records and class balance.",
        "Normalise text: lower-casing, HTML entity/URL/HTML-tag removal, repeated-character repair.",
        "Tokenise into words and clause-boundary punctuation.",
        "Mark negation scope with a neg_ prefix (the Pang et al. NOT_ trick).",
        "Remove stop-words while keeping negations and polarity connectives (but, however, very, too).",
        "Apply Porter stemming.",
        "Prepare the cleaned text for TF-IDF feature extraction.",
    ]:
        doc.add_paragraph(step, style="List Number")
    doc.add_paragraph(f"Average tokens per review after cleaning: {a['avg_tokens']}.")

    doc.add_heading("3.4 Feature Extraction", level=2)
    doc.add_paragraph(
        "Selected method: TF-IDF with unigram features built using scikit-learn TfidfVectorizer. "
        "Configuration was chosen by grid search, which also compared unigrams with unigram+bigrams. "
        "The vectoriser is fit on the training split only to avoid data leakage. sublinear_tf=True "
        "uses 1 + log(tf) so repeated words do not dominate, and L2 normalisation is applied per review.")

    doc.add_heading("3.5 Machine-Learning Models", level=2)
    doc.add_paragraph(
        "Four algorithms were compared with stratified 5-fold cross-validation on the training set "
        "using macro-F1 (the fairer metric for imbalanced data): Complement Naive Bayes, "
        "L2-regularised Logistic Regression, Linear SVM and Random Forest. The best model was then "
        "tuned with GridSearchCV over regularisation strength C, n-gram range and min_df.")

    doc.add_heading("3.6 Tools and Libraries", level=2)
    add_table(doc, ["Library / Tool", "Purpose"], [
        ["Python 3.13", "Programming language"],
        ["pandas / NumPy", "Loading, cleaning, duplicates, arrays, seeds"],
        ["scikit-learn", "TF-IDF, models, cross-validation, grid search, metrics"],
        ["NLTK (PorterStemmer, movie_reviews)", "Stemming and Part-B corpus"],
        ["matplotlib / seaborn", "Plots and confusion-matrix heatmap"],
        ["python-docx / openpyxl / nbformat", "Report, workbook and notebook generation"],
    ])

    doc.add_heading("3.7 Evaluation Results", level=2)
    doc.add_paragraph(f"Best pipeline configuration after tuning: {a['best_params']} "
                      f"(CV macro-F1 = {a['cv_f1_best']:.4f}).")
    add_table(doc, ["Metric", "Value"], [[k, f"{v:.4f}"] for k, v in tm.items()])
    doc.add_paragraph("Confusion Matrix (rows = actual, columns = predicted):")
    cm = a["confusion_matrix"]
    add_table(doc, ["Actual / Predicted", "Predicted Negative", "Predicted Positive"],
              [["Actual Negative", cm[0][0], cm[0][1]],
               ["Actual Positive", cm[1][0], cm[1][1]]])
    doc.add_paragraph("Classification Report:")
    rep = doc.add_paragraph(a["classification_report"])
    for r in rep.runs:
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
    doc.add_paragraph(
        "Because the test set contains relatively few negative reviews, a repeated 5-fold "
        "cross-validation was also run on all cleaned reviews (3 repeats). "
        "Mean results: " + ", ".join(f"{k}={v:.4f}" for k, v in a["cv_full_data"].items()) + ".")

    doc.add_heading("3.8 Error Analysis and Observations", level=2)
    for line in [
        f"The model made {a['n_errors']} errors on the test set "
        f"({100*a['n_errors']/max(a['split']['n_test'],1):.2f}%).",
        f"{a['errors_negation']} of these errors contain negation and {a['errors_mixed']} are mixed "
        "reviews that praise and complain in the same text.",
        f"{a['errors_confident_wrong']} error(s) were high-confidence (probability < 0.15 or > 0.85).",
        "Negative reviews are the hard class because the data is heavily imbalanced "
        f"({100*float(a['class_counts_clean'].get('positive',0))/max(a['dataset']['n_clean'],1):.1f}% positive).",
        "Mixed reviews and short demo sentences are difficult for a bag-of-words model.",
        f"Top positive features: {', '.join(a['top_pos'])}.",
        f"Top negative features: {', '.join(a['top_neg'])}.",
        "Improvements: collect more negative reviews, oversample/SMOTE, sentence-level scoring, "
        "word embeddings or transformer models.",
    ]:
        doc.add_paragraph(line, style="List Number")

    doc.add_heading("3.9 Demonstration Predictions", level=2)
    add_table(doc, ["Sentence", "Prediction", "P(positive)"],
              [[d["sentence"], d["prediction"], d["p_positive"]] for d in a["demo"]])

    doc.add_heading("3.10 Visualisations", level=2)
    for f, cap in [
        ("part_a_class_distribution.png", "Figure 1: Class distribution after cleaning."),
        ("part_a_model_comparison.png", "Figure 2: 5-fold cross-validation macro-F1 by model."),
        ("part_a_confusion_matrix.png", "Figure 3: Confusion matrix of the tuned Logistic Regression."),
        ("part_a_top_features.png", "Figure 4: Most positive and most negative TF-IDF features."),
    ]:
        add_image(doc, f, caption=cap)

    doc.add_page_break()
    # ------------------------------------------------------------------ PART B
    doc.add_heading("4. Part B – Mini Project: Research Paper Study and Implementation", level=1)
    doc.add_heading("4.1 Research Paper Details", level=2)
    add_table(doc, ["Item", "Details"], [
        ["Paper Title", "Thumbs up? Sentiment Classification using Machine Learning Techniques"],
        ["Authors", "Bo Pang, Lillian Lee, Shivakumar Vaithyanathan"],
        ["Venue / Year", "Proceedings of EMNLP 2002, pp. 79–86 (2002)"],
        ["URL", "https://aclanthology.org/W02-1011/ (DOI 10.3115/1118693.1118704)"],
        ["Domain", "NLP / machine learning for opinion mining and movie-review analysis"],
    ])

    doc.add_heading("4.2 Research Problem and Motivation", level=2)
    doc.add_paragraph(
        "Before this paper, most text classification sorted documents by topic. The authors ask "
        "whether sentiment, a subjective property, can be classified automatically. They use movie "
        "reviews because the star rating provides free labels, and they test whether machine learning "
        "beats simple human-chosen word lists. The main finding is that sentiment classification is "
        "harder than topic classification because opinions are subtle or negated.")

    doc.add_heading("4.3 Proposed Method / Architecture", level=2)
    for x in [
        "Corpus: 700 positive and 700 negative movie reviews.",
        "Baseline: word lists chosen by human judges (Human 1 ≈ 58%).",
        "Negation handling: words between a negation and the next punctuation mark get a NOT_ prefix.",
        "Features tested: unigrams, bigrams, unigrams+POS, adjectives only, top-N unigrams, "
        "and unigrams with position.",
        "Frequency vs presence: presence (0/1) vectors generally work better than raw counts.",
        "Classifiers: Naive Bayes, Maximum Entropy and a linear-kernel SVM.",
        "Main result: SVM on unigram presence is best (≈ 82.9%); bigrams add little.",
    ]:
        doc.add_paragraph(x, style="List Bullet")

    doc.add_heading("4.4 Dataset and Experimental Setup", level=2)
    add_table(doc, ["Component", "Pang et al. (2002)", "Our implementation"], [
        ["Dataset", "IMDb movie reviews, 700+700",
         f"NLTK movie_reviews ({b['corpus']['n_docs']} docs); 700+700 subsample across 5 seeds"],
        ["Preprocessing", "Lower-casing + NOT_ negation tagging",
         "Same: lower-casing + NOT_ tagging, no stemming/stop-words"],
        ["Models", "Naive Bayes, MaxEnt, linear SVM",
         "MultinomialNB; LogisticRegression (=MaxEnt); LinearSVC"],
        ["Hyperparameters", "Tool defaults", "NB α=1; MaxEnt C=10; SVM C=0.1"],
        ["Evaluation", "Accuracy, 3-fold CV", "Accuracy, stratified 3-fold CV × 5 subsamples"],
        ["Software", "Not described", "Python + scikit-learn (this machine)"],
    ])

    doc.add_heading("4.5 Implementation", level=2)
    doc.add_paragraph(
        "The paper's feature engineering and classifier comparison were reproduced: NOT_ negation "
        "tagging, unigram frequency, unigram presence, unigrams+bigrams presence, bigrams presence "
        "and top-2633 unigram presence, with Naive Bayes, Maximum Entropy and Linear SVM under "
        "3-fold cross-validation over reproducible 700/700 subsamples (seeds 0–4). Extras: the "
        "Human-1 word-list baseline, a negation-tagging ablation, and a 10-fold run on all 2,000 "
        "reviews. The POS, adjective-only and word-position feature sets were not implemented.")

    doc.add_heading("4.6 Results and Comparison", level=2)
    doc.add_paragraph("Our results — mean 3-fold CV accuracy (%) over 5 random 700/700 subsamples:")
    ours = pd.read_csv(os.path.join(RESULTS, "part_b_ours_mean.csv"))
    add_table(doc, list(ours.columns), ours.values.tolist())
    doc.add_paragraph("Published results — Pang et al. (2002), Table 3 (these are NOT our output):")
    paper = pd.read_csv(os.path.join(RESULTS, "part_b_paper_published.csv"))
    add_table(doc, list(paper.columns), paper.values.tolist())
    doc.add_paragraph("Difference (ours − paper), percentage points:")
    diff = pd.read_csv(os.path.join(RESULTS, "part_b_difference.csv"))
    add_table(doc, list(diff.columns), diff.values.tolist())
    doc.add_paragraph(
        f"Human-1 word-list baseline: ours {b['human_baseline']['acc']:.1f}% accuracy with "
        f"{b['human_baseline']['ties']:.1f}% ties (paper: 58% accuracy, 75% ties). "
        "Extra experiment, full 2,000 reviews, unigram presence, 10-fold CV: " +
        ", ".join(f"{k}={v:.1f}%" for k, v in b["full_corpus_10fold"].items()) + ".")

    doc.add_heading("4.7 Reproducibility and Differences", level=2)
    for x in [
        "Dataset: the paper used an earlier 1,400-review corpus; we use randomised 700/700 "
        "subsamples of the 2,000-review NLTK version, so the exact documents differ (≈ ±1 point "
        "across subsamples).",
        "Tokenisation and vocabulary cut-offs differ slightly from the paper.",
        "Tools and tuning: scikit-learn with hand-set C values replaces the original tools; this "
        "explains the larger gaps for SVM on frequency features and for MaxEnt.",
        "Not implemented: POS, adjective-only and word-position features.",
        "Cross-validation folds: the paper's exact folds are unknown, so ours are random stratified.",
        "Hardware is not reported in the paper, so no comparison is possible.",
        f"Negation ablation: {b['ablation'].get('without negation tagging', float('nan')):.1f}% "
        f"without vs {b['ablation'].get('with NOT_ negation tagging', float('nan')):.1f}% with "
        "NOT_ tagging — a marginal gain in this setup.",
    ]:
        doc.add_paragraph(x, style="List Bullet")

    doc.add_heading("4.8 Research Insights / Future Scope", level=2)
    for x in [
        "Sentiment classification is harder than topic classification (≈ 80–85% vs > 95%).",
        "Word presence beats word frequency for sentiment; bigrams add little with 1,400 documents.",
        "Unigram features are a strong baseline; negation handling gives a small gain here.",
        "Future scope: pre-trained embeddings and transformers (BERT), sentence-level analysis for "
        "mixed reviews, sarcasm and domain shift, and aspect-based sentiment analysis.",
    ]:
        doc.add_paragraph(x, style="List Bullet")
    add_image(doc, "part_b_comparison.png", caption="Figure 5: Paper vs. our replication.")
    add_image(doc, "part_b_ablation.png", caption="Figure 6: Negation-tagging ablation.")

    # ------------------------------------------------------------- conclusion
    doc.add_heading("5. Overall Conclusion", level=1)
    doc.add_paragraph(
        "Part A built a complete Amazon-review sentiment pipeline (cleaning, tokenisation, negation "
        f"marking, stop-word removal, stemming, TF-IDF and tuning) that reached {tm['accuracy']*100:.2f}% "
        f"accuracy, macro-F1 {tm['f1_macro']:.3f}, ROC-AUC {tm['roc_auc']:.3f} and MCC {tm['mcc']:.3f} "
        "on the held-out test set. Because the data is heavily imbalanced, accuracy alone is "
        "misleading, so macro-F1, MCC and the negative-class recall were reported alongside it. "
        "In Part B the main comparison of Pang et al. (2002) was reproduced on real movie reviews: "
        "unigram presence features with SVM or MaxEnt reached about 82–84%, the ranking of feature "
        "sets matched the paper, and the human word-list baseline reproduced the paper's ≈58%. "
        "Together the two parts show the full sentiment-analysis workflow and why imbalanced, "
        "real-world text is harder than clean benchmark data.")

    doc.add_heading("6. Viva Questions (selected)", level=1)
    qa = [
        ("What is sentiment analysis?",
         "It automatically identifies the opinion or polarity (positive/negative/neutral) expressed in text."),
        ("Why is preprocessing important?",
         "Raw text contains URLs, HTML, mixed case and variants that inflate the vocabulary; cleaning "
         "reduces dimensionality and exposes meaningful patterns."),
        ("What is TF-IDF?",
         "Term Frequency × Inverse Document Frequency; it weights distinctive words highly and common "
         "words lowly, separating classes better than raw counts."),
        ("Why use a test set?",
         "It gives an honest estimate on unseen data; evaluating on training data hides overfitting."),
        ("What do precision, recall and F1 indicate?",
         "Precision = share of predicted positives that are correct; recall = share of real positives "
         "found; F1 is their harmonic mean. Macro-F1 is fairer for imbalanced data."),
        ("How was the model selected?",
         "Naive Bayes, Logistic Regression, Linear SVM and Random Forest were compared with 5-fold "
         "cross-validation using macro-F1; Logistic Regression was best and was then tuned."),
        ("Main contribution of the paper?",
         "Machine-learning classifiers clearly beat human-chosen word lists for sentiment; sentiment "
         "is harder than topic classification; unigram presence features with an SVM work best."),
        ("How can the system be improved?",
         "Collect more minority-class reviews, use oversampling/SMOTE, sentence-level scoring, word "
         "embeddings or transformer models, and aspect-based analysis."),
    ]
    for q, ans in qa:
        p = doc.add_paragraph()
        p.add_run("Q: ").bold = True
        p.add_run(q)
        p2 = doc.add_paragraph()
        p2.add_run("A: ").bold = True
        p2.add_run(ans)

    doc.add_heading("7. References", level=1)
    for i, ref in enumerate([
        "Pang, B., Lee, L., & Vaithyanathan, S. (2002). Thumbs up? Sentiment Classification using "
        "Machine Learning Techniques. EMNLP 2002, pp. 79–86.",
        "Pang, B., & Lee, L. (2008). Opinion Mining and Sentiment Analysis. Foundations and Trends "
        "in Information Retrieval, 2(1–2), 1–135.",
        "Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. JMLR, 12, 2825–2830.",
        "Bird, S., Klein, E., & Loper, E. (2009). Natural Language Processing with Python. O'Reilly.",
        "Porter, M. F. (1980). An Algorithm for Suffix Stripping. Program, 14(3), 130–137.",
        "Yasser H. (n.d.). Amazon Product Reviews Dataset (Datafiniti). Kaggle. "
        "https://www.kaggle.com/datasets/yasserh/amazon-product-reviews-dataset",
    ], 1):
        doc.add_paragraph(f"[{i}] {ref}")

    doc.add_heading("8. Faculty Evaluation", level=1)
    add_table(doc, ["Evaluation Parameter", "Marks", "Remarks"], [
        ["Sentiment-analysis preprocessing", "--/5", ""],
        ["Feature extraction and model implementation", "--/5", ""],
        ["Evaluation and error analysis", "--/5", ""],
        ["Research paper understanding", "--/5", ""],
        ["Mini-project implementation", "--/5", ""],
        ["Documentation and presentation", "--/5", ""],
        ["Viva / demonstration", "--/5", ""],
        ["Total", "--/35", ""],
    ])
    doc.add_paragraph()
    doc.add_paragraph(f"Student Signature: ____________________    Date: ____________")
    doc.add_paragraph(f"Faculty Signature: ____________________")

    doc.save(DOCX)
    print("Wrote", DOCX)

    # ---- convert to PDF with LibreOffice ----
    try:
        subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", REPORT_DIR, DOCX],
            check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=240)
        pdf = DOCX[:-5] + ".pdf"
        print("Wrote", pdf if os.path.isfile(pdf) else "(PDF conversion reported success)")
    except Exception as exc:
        print(f"[warn] PDF conversion failed ({exc}); export the DOCX manually or install LibreOffice.")


if __name__ == "__main__":
    main()
