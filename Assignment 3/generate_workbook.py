#!/usr/bin/env python3
"""Build the formatted Excel workbook for Assignment 3.

The workbook is created from the actual values stored in results.json /
returned by the pipeline.  Run directly with:

    python3 generate_workbook.py
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PROJECT_DIR = Path(__file__).resolve().parent
WORKBOOK_PATH = PROJECT_DIR / "POS_Tagging_Results.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F4E79")
BODY_FONT = Font(name="Calibri", size=10)
ALT_FILL = PatternFill("solid", fgColor="EAF1F8")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

REFERENCES = [
    "Jurafsky, D., & Martin, J. H. (2025). Speech and Language Processing (3rd ed. draft). "
    "Chapters on sequence labelling and POS tagging.",
    "Manning, C. D., & Schütze, H. (1999). Foundations of Statistical Natural Language Processing. "
    "MIT Press, Chapter 10 (Part-of-Speech Tagging).",
    "Brants, T. (2000). TnT: A Statistical Part-of-Speech Tagger. Proceedings of ANLP 2000, pp. 224-231.",
    "Brill, E. (1992). A Simple Rule-Based Part of Speech Tagger. Proceedings of ANLP 1992, pp. 152-155.",
    "Ratnaparkhi, A. (1996). A Maximum Entropy Model for Part-of-Speech Tagging. Proceedings of EMNLP 1996.",
    "Lafferty, J., McCallum, A., & Pereira, F. (2001). Conditional Random Fields: Probabilistic Models "
    "for Segmenting and Labeling Sequence Data. Proceedings of ICML 2001.",
    "Huang, Z., Xu, W., & Yu, K. (2015). Bidirectional LSTM-CRF Models for Sequence Tagging. "
    "arXiv:1508.01991.",
    "Devlin, J., Chang, M.-W., Lee, K., & Toutanova, K. (2019). BERT: Pre-training of Deep Bidirectional "
    "Transformers for Language Understanding. Proceedings of NAACL-HLT 2019.",
    "Marcus, M. P., Santorini, B., & Marcinkiewicz, M. A. (1993). Building a Large Annotated Corpus of "
    "English: The Penn Treebank. Computational Linguistics, 19(2), 313-330.",
    "Bird, S., Klein, E., & Loper, E. (2009). Natural Language Processing with Python. O'Reilly Media.",
    "NLTK documentation and corpora: https://www.nltk.org/",
    "Universal Dependencies POS tags: https://universaldependencies.org/u/pos/",
]


def _style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def _autosize(ws, max_width: int = 60) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        longest = 0
        for cell in column_cells:
            if cell.value is None:
                continue
            for line in str(cell.value).splitlines():
                longest = max(longest, len(line))
        ws.column_dimensions[letter].width = min(max(10, longest + 2), max_width)


def _write_table(ws, headers: List[str], rows: List[list], freeze: str = "A2", widths=None) -> None:
    ws.append(headers)
    _style_header(ws, 1)
    for row in rows:
        ws.append(row)
    for index, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP
            if index % 2 == 0:
                cell.fill = ALT_FILL
    ws.freeze_panes = freeze
    if widths:
        for letter, width in widths.items():
            ws.column_dimensions[letter].width = width
    else:
        _autosize(ws)


def _write_key_value(ws, rows: List[tuple], value_format: str | None = None) -> None:
    ws.append(["Metric", "Value"])
    _style_header(ws, 1)
    for key, value in rows:
        ws.append([key, value])
    for index, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP
            if index % 2 == 0:
                cell.fill = ALT_FILL
        if value_format and isinstance(row[1].value, (int, float)) and not isinstance(row[1].value, bool):
            row[1].number_format = value_format
    ws.freeze_panes = "A2"
    _autosize(ws, max_width=100)


def build_workbook(result: Dict) -> Path:
    """Create and save the nine-sheet Excel workbook."""
    corpus = result["corpus"]
    hmm = result["hmm"]
    baseline = result["baseline"]
    config = result["config"]

    wb = Workbook()
    wb.properties.title = "NLP Laboratory Assignment 3 - POS Tagging Results"
    wb.properties.creator = result.get("student_name", "Shravani Kurkute")
    wb.properties.subject = "HMM POS tagging with Viterbi decoding on the Penn Treebank"

    # 1. Summary -------------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws.append(["NLP Laboratory Assignment 3 - Survey and Implementation of POS Tagging"])
    ws.merge_cells("A1:B1")
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.append(["Student", result.get("student_name", "Shravani Kurkute")])
    ws.append(["Generated", result.get("generated_at", "")])
    ws.append([])
    summary_rows = [
        ("Corpus", corpus["name"]),
        ("Universal tagset", config["tagset"]),
        ("Prepared sentences", corpus["prepared_sentences"]),
        ("Training sentences (80%)", corpus["train_sentences"]),
        ("Test sentences (20%)", corpus["test_sentences"]),
        ("Training tokens", corpus["train_tokens"]),
        ("Test tokens", corpus["test_tokens"]),
        ("Training vocabulary", corpus["train_vocabulary"]),
        ("Test OOV tokens", corpus["test_oov_tokens"]),
        ("Test OOV rate", corpus["test_oov_rate"]),
        ("Baseline accuracy", baseline["accuracy"]),
        ("HMM accuracy", hmm["accuracy"]),
        ("HMM macro precision", hmm["macro_precision"]),
        ("HMM macro recall", hmm["macro_recall"]),
        ("HMM macro F1", hmm["macro_f1"]),
        ("HMM weighted precision", hmm["weighted_precision"]),
        ("HMM weighted recall", hmm["weighted_recall"]),
        ("HMM weighted F1", hmm["weighted_f1"]),
        ("HMM OOV accuracy", hmm["oov_accuracy"]),
        ("HMM known-token accuracy", hmm["known_accuracy"]),
        ("HMM total token errors", result["total_error_tokens"]),
        ("Random seed", config["seed"]),
        ("Transition smoothing k", config["transition_k"]),
        ("Pipeline runtime (seconds)", result.get("runtime_seconds", "")),
    ]
    ws.append(["Metric", "Value"])
    _style_header(ws, 5)
    start_row = 6
    for offset, (key, value) in enumerate(summary_rows):
        row = start_row + offset
        ws.cell(row=row, column=1, value=key)
        ws.cell(row=row, column=2, value=value)
        for column in (1, 2):
            cell = ws.cell(row=row, column=column)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP
            if row % 2 == 0:
                cell.fill = ALT_FILL
        if isinstance(value, float):
            ws.cell(row=row, column=2).number_format = "0.000000"
        elif isinstance(value, int):
            ws.cell(row=row, column=2).number_format = "#,##0"
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 78
    ws.freeze_panes = "A6"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "1F4E79"

    # 2. Dataset_Statistics --------------------------------------------------
    ws = wb.create_sheet("Dataset_Statistics")
    dataset_rows = [
        ("Corpus", corpus["name"]),
        ("Raw Penn Treebank sentences", corpus["raw_sentences"]),
        ("Raw tokens", corpus["raw_tokens"]),
        ("Trace (-NONE-) tokens removed", corpus["trace_tokens_removed"]),
        ("Sentences after trace removal", corpus["sentences_after_trace_removal"]),
        ("Tokens after trace removal", corpus["tokens_after_trace_removal"]),
        ("Exact duplicate sentences removed", corpus["duplicate_sentences_removed"]),
        ("Prepared sentences", corpus["prepared_sentences"]),
        ("Prepared tokens", corpus["prepared_tokens"]),
        ("Training sentences", corpus["train_sentences"]),
        ("Test sentences", corpus["test_sentences"]),
        ("Training tokens", corpus["train_tokens"]),
        ("Test tokens", corpus["test_tokens"]),
        ("Training vocabulary size", corpus["train_vocabulary"]),
        ("Test OOV tokens", corpus["test_oov_tokens"]),
        ("Test OOV rate", corpus["test_oov_rate"]),
        ("Number of PTB tags mapped", corpus["ptb_tag_count"]),
        ("Number of Universal tags", corpus["universal_tag_count"]),
        ("Sentence overlap between splits", 0),
        ("Random seed", config["seed"]),
        ("Test fraction", config["test_size"]),
    ]
    _write_key_value(ws, dataset_rows)
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=2)
        if isinstance(cell.value, float):
            cell.number_format = "0.0000"
        elif isinstance(cell.value, int):
            cell.number_format = "#,##0"
    ws.sheet_properties.tabColor = "2E75B6"

    # 3. Training_Configuration ---------------------------------------------
    ws = wb.create_sheet("Training_Configuration")
    configuration_rows = [
        ("Model", "First-order Hidden Markov Model (supervised)"),
        ("Decoder", config["decoder"]),
        ("Baseline", config["baseline"]),
        ("Tagset", config["tagset"]),
        ("Random seed", config["seed"]),
        ("Train/test split", "80:20 sentence-level"),
        ("Test fraction", config["test_size"]),
        ("Transition smoothing", config["transition_smoothing"]),
        ("Transition smoothing k", config["transition_k"]),
        ("Emission probabilities", "Maximum likelihood for observed words"),
        ("Rare-word threshold (frequency <=)", config["rare_word_threshold"]),
        ("Maximum suffix length", config["max_suffix_length"]),
        ("OOV add-alpha smoothing", config["oov_alpha"]),
        ("OOV numeric handling", "Digit-shape prior learned from rare numeric tokens"),
        ("OOV suffix handling", "Longest matching suffix from rare training words"),
        ("Unknown-word fallback", "Tag prior + shape prior; all tag scores positive"),
        ("Vocabulary size", result["model"]["vocabulary_size"]),
        ("Emission pairs (tag, word)", result["model"]["emission_pairs"]),
        ("Rare-word occurrences used for OOV model", result["model"]["rare_token_occurrences"]),
        ("Numeric rare-token occurrences", result["model"]["numeric_rare_occurrences"]),
        ("Python version", result["environment"].get("python", "")),
        ("NLTK version", result["environment"].get("nltk", "")),
        ("NumPy version", result["environment"].get("numpy", "")),
        ("scikit-learn version", result["environment"].get("scikit-learn", "")),
        ("Pandas version", result["environment"].get("pandas", "")),
        ("Matplotlib version", result["environment"].get("matplotlib", "")),
    ]
    _write_key_value(ws, configuration_rows)
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=2)
        if isinstance(cell.value, int) and not isinstance(cell.value, bool):
            cell.number_format = "#,##0"
        elif isinstance(cell.value, float):
            cell.number_format = "0.0000"
    ws.sheet_properties.tabColor = "70AD47"

    # 4. Baseline_vs_HMM -----------------------------------------------------
    ws = wb.create_sheet("Baseline_vs_HMM")
    metric_rows = [
        ("Accuracy", baseline["accuracy"], hmm["accuracy"]),
        ("Macro precision", baseline["macro_precision"], hmm["macro_precision"]),
        ("Macro recall", baseline["macro_recall"], hmm["macro_recall"]),
        ("Macro F1", baseline["macro_f1"], hmm["macro_f1"]),
        ("Weighted precision", baseline["weighted_precision"], hmm["weighted_precision"]),
        ("Weighted recall", baseline["weighted_recall"], hmm["weighted_recall"]),
        ("Weighted F1", baseline["weighted_f1"], hmm["weighted_f1"]),
        ("OOV accuracy", baseline["oov_accuracy"], hmm["oov_accuracy"]),
        ("Known-token accuracy", baseline["known_accuracy"], hmm["known_accuracy"]),
    ]
    rows = [[name, base, hmm_value, hmm_value - base] for name, base, hmm_value in metric_rows]
    _write_table(
        ws,
        ["Metric", "Baseline", "HMM", "Difference (HMM - Baseline)"],
        rows,
        widths={"A": 28, "B": 14, "C": 14, "D": 28},
    )
    for row in range(2, ws.max_row + 1):
        for column in range(2, 5):
            ws.cell(row=row, column=column).number_format = "0.000000"
    ws.auto_filter.ref = f"A1:D{ws.max_row}"
    ws.sheet_properties.tabColor = "ED7D31"

    # 5. Per_Tag_Metrics -----------------------------------------------------
    ws = wb.create_sheet("Per_Tag_Metrics")
    tags = hmm["per_tag"].keys()
    rows = []
    for tag in tags:
        b = result["baseline_per_tag"].get(tag, {})
        h = hmm["per_tag"].get(tag, {})
        rows.append(
            [
                tag,
                b.get("precision", 0.0),
                b.get("recall", 0.0),
                b.get("f1_score", 0.0),
                b.get("support", 0),
                h.get("precision", 0.0),
                h.get("recall", 0.0),
                h.get("f1_score", 0.0),
                h.get("support", 0),
                h.get("f1_score", 0.0) - b.get("f1_score", 0.0),
            ]
        )
    _write_table(
        ws,
        [
            "Tag",
            "Baseline Precision",
            "Baseline Recall",
            "Baseline F1",
            "Baseline Support",
            "HMM Precision",
            "HMM Recall",
            "HMM F1",
            "HMM Support",
            "F1 Difference",
        ],
        rows,
        widths={"A": 9, "B": 14, "C": 14, "D": 14, "E": 12, "F": 14, "G": 14, "H": 14, "I": 12, "J": 14},
    )
    for row in range(2, ws.max_row + 1):
        for column in range(2, 11):
            if column in (5, 9):
                ws.cell(row=row, column=column).number_format = "#,##0"
            else:
                ws.cell(row=row, column=column).number_format = "0.0000"
    ws.auto_filter.ref = f"A1:J{ws.max_row}"
    ws.sheet_properties.tabColor = "A5A5A5"

    # 6. Sample_Tagged_Output ------------------------------------------------
    ws = wb.create_sheet("Sample_Tagged_Output")
    sample_rows = result.get("sample_tagged_sentences", [])
    rows = [
        [
            row["sentence_id"],
            row["token_index"],
            row["token"],
            row["gold_tag"],
            row["hmm_predicted_tag"],
            row["baseline_predicted_tag"],
            "Yes" if row["is_oov"] else "No",
        ]
        for row in sample_rows
    ]
    _write_table(
        ws,
        ["Sentence ID", "Token Index", "Token", "Gold Tag", "HMM Prediction", "Baseline Prediction", "OOV"],
        rows,
        widths={"A": 11, "B": 11, "C": 24, "D": 11, "E": 15, "F": 18, "G": 8},
    )
    for row in range(2, ws.max_row + 1):
        for column in (1, 2):
            ws.cell(row=row, column=column).number_format = "#,##0"
        ws.cell(row=row, column=3).alignment = Alignment(horizontal="left", vertical="top")
    ws.auto_filter.ref = f"A1:G{ws.max_row}"
    ws.sheet_properties.tabColor = "7030A0"

    # 7. Error_Analysis ------------------------------------------------------
    ws = wb.create_sheet("Error_Analysis")
    error_rows = []
    for rank, error in enumerate(result.get("top_errors", []), start=1):
        error_rows.append(
            [
                rank,
                error["gold_tag"],
                error["predicted_tag"],
                error["count"],
                error["percentage_of_errors"] / 100.0,
                error["example_context"],
            ]
        )
    _write_table(
        ws,
        ["Rank", "Gold Tag", "Predicted Tag", "Count", "% of All Errors", "Example Context"],
        error_rows,
        widths={"A": 7, "B": 11, "C": 13, "D": 9, "E": 15, "F": 90},
    )
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=4).number_format = "#,##0"
        ws.cell(row=row, column=5).number_format = "0.00%"
        ws.cell(row=row, column=6).alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    ws.auto_filter.ref = f"A1:F{ws.max_row}"
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = "C00000"

    # 8. OOV_Analysis --------------------------------------------------------
    ws = wb.create_sheet("OOV_Analysis")
    oov_rows = [
        ("OOV definition", "Token word type absent from the training vocabulary"),
        ("Test tokens", hmm["support"]),
        ("OOV tokens", hmm["oov_count"]),
        ("OOV rate", hmm["oov_rate"]),
        ("Known tokens", hmm["known_count"]),
        ("HMM OOV accuracy", hmm["oov_accuracy"]),
        ("HMM known-token accuracy", hmm["known_accuracy"]),
        ("Baseline OOV accuracy", baseline["oov_accuracy"]),
        ("Baseline known-token accuracy", baseline["known_accuracy"]),
        ("OOV accuracy gain (HMM - Baseline)", hmm["oov_accuracy"] - baseline["oov_accuracy"]),
        ("Rare-word threshold", config["rare_word_threshold"]),
        ("Maximum suffix length", config["max_suffix_length"]),
        ("OOV add-alpha", config["oov_alpha"]),
        ("Numeric strategy", "Numeric-shape prior from rare numeric training tokens"),
        ("Suffix strategy", "Longest matching suffix (1-4 characters) from rare training words"),
    ]
    _write_key_value(ws, oov_rows)
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=2)
        if isinstance(cell.value, float):
            cell.number_format = "0.000000" if not cell.value <= 1 else "0.00%"
        elif isinstance(cell.value, int):
            cell.number_format = "#,##0"
    ws.sheet_properties.tabColor = "FFC000"

    # 9. References ----------------------------------------------------------
    ws = wb.create_sheet("References")
    rows = [[index, reference] for index, reference in enumerate(REFERENCES, start=1)]
    _write_table(ws, ["No.", "Reference"], rows, widths={"A": 6, "B": 150})
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.sheet_properties.tabColor = "00B0F0"

    for sheet in wb.worksheets:
        sheet.sheet_view.zoomScale = 100
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(WORKBOOK_PATH)

    # Validate that the saved workbook can be reopened (LibreOffice-compatible).
    check = load_workbook(WORKBOOK_PATH, read_only=True)
    expected = [
        "Summary",
        "Dataset_Statistics",
        "Training_Configuration",
        "Baseline_vs_HMM",
        "Per_Tag_Metrics",
        "Sample_Tagged_Output",
        "Error_Analysis",
        "OOV_Analysis",
        "References",
    ]
    if check.sheetnames != expected:
        raise RuntimeError(f"Workbook sheets are incorrect: {check.sheetnames}")
    check.close()
    return WORKBOOK_PATH


def main() -> int:
    with (PROJECT_DIR / "results.json").open(encoding="utf-8") as handle:
        result = json.load(handle)
    path = build_workbook(result)
    print(f"Workbook created: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
