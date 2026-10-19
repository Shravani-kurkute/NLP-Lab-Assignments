"""Automated tests for the HMM POS-tagging assignment.

Run from the project directory:
    python3 -m pytest tests -v

The tests use the same corpus, split, seed and configuration as the main
pipeline.  They train a fresh model once per test session (about 2-3 seconds).
"""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pytest

PROJECT_DIR = Path(__file__).resolve().parents[1]
import sys

if str(PROJECT_DIR) not in sys.path:
    sys.path.insert(0, str(PROJECT_DIR))

from pos_tagging_hmm import (  # noqa: E402
    DEFAULT_CUSTOM_SENTENCES,
    HMMPosTagger,
    MostFrequentTagBaseline,
    SEED,
    TRANSITION_K,
    UNIVERSAL_TAG_DESCRIPTIONS,
    deduplicate_sentences,
    ensure_nltk_resources,
    evaluate_tags,
    flatten_predictions,
    flatten_tagged,
    load_custom_sentences,
    load_treebank_sentences,
    map_ptb_to_universal,
    prepare_dataset,
    tokenize_sentence,
    write_sentence_file,
)


@pytest.fixture(scope="session")
def prepared():
    ensure_nltk_resources(allow_download=True, verbose=False)
    raw, _ = load_treebank_sentences(remove_traces=True)
    mapped, mapping = map_ptb_to_universal(raw)
    unique, _ = deduplicate_sentences(mapped)
    train, test, meta = prepare_dataset(unique, test_size=0.20, seed=SEED)
    return {
        "train": train,
        "test": test,
        "meta": meta,
        "mapping": mapping,
        "unique": unique,
    }


@pytest.fixture(scope="session")
def models(prepared):
    hmm = HMMPosTagger(
        transition_k=TRANSITION_K,
        rare_word_threshold=2,
        max_suffix_len=4,
        oov_alpha=0.1,
    ).train(prepared["train"])
    baseline = MostFrequentTagBaseline().train(prepared["train"])
    return hmm, baseline


# ---------------------------------------------------------------------------
# Corpus, tagset and split tests
# ---------------------------------------------------------------------------


def test_trace_tokens_are_removed():
    raw, stats = load_treebank_sentences(remove_traces=True)
    assert stats["trace_tokens_removed"] > 0
    assert all(tag != "-NONE-" for sent in raw for _, tag in sent)


def test_universal_mapping_and_tagset(prepared):
    mapping = prepared["mapping"]
    assert "-NONE-" not in mapping or mapping["-NONE-"] == "X"
    assert set(mapping.values()).issubset(set(UNIVERSAL_TAG_DESCRIPTIONS))
    assert mapping["NN"] == "NOUN"
    assert mapping["VBZ"] == "VERB"
    assert mapping["."] == "."


def test_reproducible_split(prepared):
    train2, test2, meta2 = prepare_dataset(prepared["unique"], test_size=0.20, seed=SEED)
    assert train2 == prepared["train"]
    assert test2 == prepared["test"]
    assert meta2["train_sentences"] == prepared["meta"]["train_sentences"]


def test_different_seed_changes_split(prepared):
    train2, _, _ = prepare_dataset(prepared["unique"], test_size=0.20, seed=7)
    assert train2 != prepared["train"]


def test_no_sentence_overlap(prepared):
    train_keys = {tuple(s) for s in prepared["train"]}
    test_keys = {tuple(s) for s in prepared["test"]}
    assert train_keys.isdisjoint(test_keys)
    assert len(train_keys) + len(test_keys) == len(prepared["unique"])


def test_split_sizes(prepared):
    assert prepared["meta"]["train_sentences"] == 3122
    assert prepared["meta"]["test_sentences"] == 781
    assert prepared["meta"]["sentence_overlap"] == 0


# ---------------------------------------------------------------------------
# HMM and Viterbi tests
# ---------------------------------------------------------------------------


def test_hmm_training_statistics(models):
    hmm, _ = models
    assert hmm.is_trained
    assert len(hmm.tags) == 12
    assert len(hmm.vocab) == 10625
    assert all(math_is_finite(row[tag]) for row in hmm.log_transition.values() for tag in row)


def math_is_finite(value):
    return isinstance(value, float) and np.isfinite(value)


def test_empty_input_is_safe(models):
    hmm, baseline = models
    assert hmm.tag([]) == []
    assert baseline.tag([]) == []
    tags, log_prob, details = hmm.viterbi_decode([], return_details=True)
    assert tags == []
    assert log_prob == 0.0
    assert details["score_matrix"] == []


def test_single_word_sentence(models):
    hmm, baseline = models
    tags = hmm.tag(["dog"])
    assert len(tags) == 1
    assert tags[0] in hmm.tags
    assert len(baseline.tag(["dog"])) == 1


def test_multiword_sentence(models):
    hmm, _ = models
    tokens = tokenize_sentence("The quick brown fox jumps over the lazy dog.")
    tags = hmm.tag(tokens)
    assert len(tags) == len(tokens)
    assert all(tag in hmm.tags for tag in tags)


def test_viterbi_deterministic_and_negative_logprob(models):
    hmm, _ = models
    tokens = tokenize_sentence("The committee approved the budget after a long discussion.")
    tags1, score1 = hmm.viterbi_decode(tokens)
    tags2, score2 = hmm.viterbi_decode(tokens)
    assert tags1 == tags2
    assert score1 == pytest.approx(score2)
    assert score1 <= 0.0


def test_viterbi_details_shapes(models):
    hmm, _ = models
    tokens = ["The", "company", "reported", "growth", "."]
    tags, score, details = hmm.viterbi_decode(tokens, return_details=True)
    assert len(tags) == len(tokens)
    assert len(details["score_matrix"]) == len(tokens)
    assert all(len(row) == len(hmm.tags) for row in details["score_matrix"])
    assert all(len(row) == len(hmm.tags) for row in details["backpointers"])


def test_unknown_word_handling(models):
    hmm, _ = models
    for token in ["Shravanikurkute", "zzzqqq", "frobnicating"]:
        probs = hmm.oov_emission_probs(token)
        assert set(probs) == set(hmm.tags)
        assert sum(probs.values()) == pytest.approx(1.0)
        assert all(value > 0 for value in probs.values())
    tags = hmm.tag(tokenize_sentence("Shravanikurkute visited the new laboratory."))
    assert len(tags) == 6
    assert all(tag in hmm.tags for tag in tags)


def test_numeric_tokens_receive_num(models):
    hmm, _ = models
    tokens = tokenize_sentence("The company earned 12345 dollars in 2025.")
    tags = hmm.tag(tokens)
    assert tags[tokens.index("12345")] == "NUM"
    assert tags[tokens.index("2025")] == "NUM"


def test_smoothed_transitions_are_positive(models):
    hmm, _ = models
    for prev in [hmm.START] + hmm.tags:
        for curr in hmm.tags + [hmm.END]:
            assert hmm.log_transition[prev][curr] > float("-inf")
            assert np.isfinite(hmm.log_transition[prev][curr])


# ---------------------------------------------------------------------------
# Evaluation tests
# ---------------------------------------------------------------------------


def test_evaluation_metrics_calculation():
    gold = ["NOUN", "VERB", "NOUN", "ADJ"]
    pred = ["NOUN", "NOUN", "NOUN", "ADJ"]
    words = ["a", "b", "c", "d"]
    labels = ["ADJ", "NOUN", "VERB", "."]
    metrics, per_tag, cm = evaluate_tags(gold, pred, labels, words, set(words))
    assert metrics["accuracy"] == pytest.approx(0.75)
    assert per_tag.loc[per_tag["tag"] == "VERB", "recall"].iloc[0] == pytest.approx(0.0)
    assert cm.shape == (4, 4)
    assert metrics["oov_count"] == 0


def test_oov_metric_calculation():
    gold = ["NOUN", "VERB"]
    pred = ["NOUN", "VERB"]
    words = ["known", "unseenword"]
    labels = ["NOUN", "VERB"]
    metrics, _, _ = evaluate_tags(gold, pred, labels, words, {"known"})
    assert metrics["oov_count"] == 1
    assert metrics["oov_accuracy"] == pytest.approx(1.0)
    assert metrics["known_accuracy"] == pytest.approx(1.0)


def test_baseline_and_hmm_use_same_test_tokens(prepared, models):
    hmm, baseline = models
    test_words, test_tags = flatten_tagged(prepared["test"])
    hmm_preds = [hmm.tag([w for w, _ in sent]) for sent in prepared["test"]]
    base_preds = [baseline.tag([w for w, _ in sent]) for sent in prepared["test"]]
    assert len(flatten_predictions(hmm_preds)) == len(test_words)
    assert len(flatten_predictions(base_preds)) == len(test_words)
    assert len(test_tags) == len(test_words)


def test_hmm_outperforms_baseline_on_test(prepared, models):
    hmm, baseline = models
    test_words, y_true = flatten_tagged(prepared["test"])
    train_words, _ = flatten_tagged(prepared["train"])
    vocab = set(train_words)
    hmm_pred = flatten_predictions([hmm.tag([w for w, _ in s]) for s in prepared["test"]])
    base_pred = flatten_predictions([baseline.tag([w for w, _ in s]) for s in prepared["test"]])
    hmm_metrics, _, _ = evaluate_tags(y_true, hmm_pred, hmm.tags, test_words, vocab)
    base_metrics, _, _ = evaluate_tags(y_true, base_pred, hmm.tags, test_words, vocab)
    assert hmm_metrics["accuracy"] > base_metrics["accuracy"]
    assert hmm_metrics["oov_accuracy"] > 0.0


# ---------------------------------------------------------------------------
# Custom sentences and output files
# ---------------------------------------------------------------------------


def test_custom_sentence_file_has_at_least_twenty():
    sentences = load_custom_sentences(PROJECT_DIR / "custom_test_sentences.txt")
    assert len(sentences) >= 20
    assert len(DEFAULT_CUSTOM_SENTENCES) >= 20


def test_sentence_file_writer(tmp_path):
    path = tmp_path / "sentences.txt"
    write_sentence_file(path, [[("Hello", "X"), ("world", "NOUN")], [(".", ".")]])
    lines = path.read_text(encoding="utf-8").splitlines()
    assert lines == ["Hello world", "."]


def test_required_output_files_exist():
    required = [
        "results.json",
        "evaluation_results.csv",
        "sample_tagged_output.csv",
        "error_analysis.csv",
        "train_sentences.txt",
        "test_sentences.txt",
        "custom_test_sentences.txt",
        "visualizations/confusion_matrix.png",
        "visualizations/per_tag_scores.png",
        "visualizations/model_comparison.png",
        "visualizations/oov_analysis.png",
        "visualizations/tag_distribution.png",
    ]
    for relative in required:
        assert (PROJECT_DIR / relative).exists(), f"Missing generated artefact: {relative}"


def test_results_json_is_valid():
    path = PROJECT_DIR / "results.json"
    with path.open(encoding="utf-8") as handle:
        result = json.load(handle)
    assert result["config"]["seed"] == 42
    assert result["hmm"]["accuracy"] > 0.9
    assert result["corpus"]["train_sentences"] > 0
    assert result["corpus"]["test_sentences"] > 0
    assert len(result["custom_sentences"]) >= 20


def test_workbook_and_report_if_present():
    workbook_path = PROJECT_DIR / "POS_Tagging_Results.xlsx"
    report_path = PROJECT_DIR / "POS_Tagging_Assignment_03_Report.docx"
    if workbook_path.exists():
        from openpyxl import load_workbook

        workbook = load_workbook(workbook_path, read_only=True)
        assert "Summary" in workbook.sheetnames
        assert "Per_Tag_Metrics" in workbook.sheetnames
        workbook.close()
    if report_path.exists():
        from docx import Document

        document = Document(report_path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert "Shravani Kurkute" in text
        assert "Viterbi" in text
