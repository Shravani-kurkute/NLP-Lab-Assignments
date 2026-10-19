#!/usr/bin/env python3
"""Build results/Assignment4_Results.xlsx from the evaluation outputs.

Sheets:
    Overview, Accuracy Summary, Filtered vs Full, Test15 Predictions,
    Large1000 Predictions, Test15 Errors, Large1000 Errors, Error Breakdown
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parents[1]
RESULTS_DIR = BASE_DIR / "results"
WORKBOOK_PATH = RESULTS_DIR / "Assignment4_Results.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
CORRECT_FILL = PatternFill("solid", fgColor="C6EFCE")
INCORRECT_FILL = PatternFill("solid", fgColor="FFC7CE")
SECTION_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def read_csv(name: str) -> pd.DataFrame:
    return pd.read_csv(RESULTS_DIR / name)


def overview_frame(summary: dict) -> pd.DataFrame:
    rows = [
        ("Assignment", "NLP Laboratory Assignment 04 - Country Prediction using Word Embeddings"),
        ("Student name", "Shravani Balasaheb Kurkute"),
        ("Roll number", "421129"),
        ("PRN / Registration no.", "22311926"),
        ("Class / Division", "BTechA"),
        ("Academic year", "2026-2027"),
        ("Faculty name", "Dr. Kapil N. Vhatkar"),
        ("Embedding model", summary["model"]["name"]),
        ("Embedding source", summary["model"]["source"]),
        ("Vocabulary size", f"{summary['model']['vocabulary_size']:,} words"),
        ("Embedding dimension", summary["model"]["dimension"]),
        ("Candidate countries", summary["candidate_countries"]),
        ("Valid capital-country pairs", 78),
        ("Test analogies", summary["datasets"]["test15"]["size"]),
        ("Evaluation analogies", summary["datasets"]["large1000"]["size"]),
        ("Analogy formula", "vector(country1) - vector(capital1) + vector(capital2)"),
        ("Similarity measure", "Cosine similarity on L2-normalised vectors"),
        ("Generated at", summary["generated_at"]),
        ("Python version", summary["environment"]["python"]),
        ("Platform", summary["environment"]["platform"]),
    ]
    return pd.DataFrame(rows, columns=["Item", "Value"])


def accuracy_frame(summary: dict) -> pd.DataFrame:
    rows = []
    for dataset, modes in summary["results"].items():
        for mode, metrics in modes.items():
            rows.append(
                {
                    "dataset": dataset,
                    "mode": mode,
                    "total": metrics["total"],
                    "correct_top1": metrics["correct_top1"],
                    "incorrect_top1": metrics["incorrect_top1"],
                    "top1_accuracy_pct": metrics["top1_accuracy_pct"],
                    "correct_top5": metrics["correct_top5"],
                    "top5_accuracy_pct": metrics["top5_accuracy_pct"],
                    "mean_predicted_score": metrics["mean_predicted_score"],
                    "mean_actual_score": metrics["mean_actual_score"],
                    "mean_top1_margin": metrics["mean_top1_margin"],
                    "median_top1_margin": metrics["median_top1_margin"],
                    "runtime_seconds": metrics["runtime_seconds"],
                }
            )
    return pd.DataFrame(rows)


def comparison_frame(summary: dict) -> pd.DataFrame:
    rows = []
    for dataset, metrics in summary["comparison"].items():
        filtered = summary["results"][dataset]["country_filtered"]
        full = summary["results"][dataset]["full_vocabulary"]
        rows.append(
            {
                "dataset": dataset,
                "country_filtered_top1_pct": filtered["top1_accuracy_pct"],
                "full_vocabulary_top1_pct": full["top1_accuracy_pct"],
                "top1_gain_pct_points": metrics["top1_gain_country_filtered_pct_points"],
                "country_filtered_top5_pct": filtered["top5_accuracy_pct"],
                "full_vocabulary_top5_pct": full["top5_accuracy_pct"],
                "top5_gain_pct_points": metrics["top5_gain_country_filtered_pct_points"],
                "errors_country_filtered": filtered["incorrect_top1"],
                "errors_full_vocabulary": full["incorrect_top1"],
            }
        )
    return pd.DataFrame(rows)


def breakdown_frame(summary: dict) -> pd.DataFrame:
    return pd.DataFrame(summary.get("error_breakdown", []), columns=["country", "errors"])


def style_sheet(ws, highlight_correct: bool = False) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    correct_column = None
    headers = [cell.value for cell in ws[1]]
    if highlight_correct and "correct" in headers:
        correct_column = headers.index("correct") + 1

    for row in ws.iter_rows(min_row=1):
        for cell in row:
            cell.border = BORDER
            if cell.row > 1:
                cell.alignment = Alignment(vertical="top", wrap_text=False)

    if correct_column:
        for row_index in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_index, column=correct_column)
            if cell.value == "Correct":
                cell.fill = CORRECT_FILL
                cell.font = Font(color="006100", bold=True)
            elif cell.value == "Incorrect":
                cell.fill = INCORRECT_FILL
                cell.font = Font(color="9C0006", bold=True)

    for index in range(1, ws.max_column + 1):
        letter = get_column_letter(index)
        longest = max(
            (len(str(ws.cell(row=row, column=index).value or ""))
             for row in range(1, min(ws.max_row, 400) + 1)),
            default=10,
        )
        ws.column_dimensions[letter].width = min(max(10, longest + 2), 55)


def main() -> None:
    summary = json.loads((RESULTS_DIR / "accuracy_summary.json").read_text())

    frames = {
        "Overview": overview_frame(summary),
        "Accuracy Summary": accuracy_frame(summary),
        "Filtered vs Full": comparison_frame(summary),
        "Test15 Predictions": read_csv("predictions_test15_country_filtered.csv"),
        "Large1000 Predictions": read_csv("predictions_large1000_country_filtered.csv"),
        "Test15 Errors": read_csv("error_analysis_test15_country_filtered.csv"),
        "Large1000 Errors": read_csv("error_analysis_large1000_country_filtered.csv"),
        "Error Breakdown": breakdown_frame(summary),
    }

    WORKBOOK_PATH.parent.mkdir(exist_ok=True)
    with pd.ExcelWriter(WORKBOOK_PATH, engine="openpyxl") as writer:
        for name, frame in frames.items():
            frame.to_excel(writer, sheet_name=name, index=False)

    from openpyxl import load_workbook

    workbook = load_workbook(WORKBOOK_PATH)
    for name in workbook.sheetnames:
        sheet = workbook[name]
        style_sheet(sheet, highlight_correct=name.endswith("Predictions"))
        if name == "Overview":
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row=row, column=1).fill = SECTION_FILL
                sheet.cell(row=row, column=1).font = Font(bold=True)
    workbook.save(WORKBOOK_PATH)

    print(f"workbook written to {WORKBOOK_PATH}")
    for name, frame in frames.items():
        print(f"  {name:<22} {len(frame):>5} rows x {len(frame.columns)} cols")


if __name__ == "__main__":
    main()
